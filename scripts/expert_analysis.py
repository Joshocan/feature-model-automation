#!/usr/bin/env python3
"""Reproduce the paper's expert assessment from immutable saved inputs.

No generation, embeddings, workbook edits, or personal metadata exports. A failed
published-value check produces only a failure report, never final analysis tables.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
import hashlib
from importlib.metadata import version
from itertools import combinations
import math
from pathlib import Path
import statistics
import sys
import xml.etree.ElementTree as ET

import numpy as np
from openpyxl import load_workbook
from scipy.stats import kendalltau, rankdata

try:
    import krippendorff
except ImportError:
    raise SystemExit("Install the expert-analysis dependency: python -m pip install krippendorff==0.8.2")

ROOT = Path(__file__).resolve().parents[1]
RATERS = ("Expert A", "Expert B")
CORPORA = ("repair", "federation")
POLICIES = ("independent_max", "one_to_one")
DIMENSIONS = ("Q1", "Q2", "Q3", "Q4", "Q5")
DECISIONS = {"Reject": 0, "Accept with revision": 1, "Accept": 2, "Cannot judge": None}
CITATIONS = {"Yes": 2, "Partially": 1, "No": 0, "Cannot judge": None}
B_CITATION_MAP = {"OK": "Yes", "Not exactly": "Partially",
                  "not sure, depends on the context": "Cannot judge", "No": "No"}
MODELS = {"deepseek-v4.1-flash:cloud": "DeepSeek", "glm-5.3-flash:cloud": "GLM",
          "gpt-6-astra": "Astra"}
FEATURE_TAGS = {"and", "or", "alt", "feature"}


class SanityError(ValueError):
    """A supplied paper check does not reproduce."""


class Checks:
    def __init__(self):
        self.rows = []

    def equal(self, name, actual, expected, decimals=None):
        # Half a unit in the last published decimal, not an arbitrary tolerance.
        ok = (actual == expected if decimals is None else
              actual is not None and math.isfinite(actual)
              and abs(actual - expected) <= 0.5 * 10 ** -decimals + 1e-12)
        self.rows.append(dict(check=name, observed=actual, expected=expected,
                              decimal_places=decimals, passed=ok))
        if not ok:
            raise SanityError(f"{name}: observed {actual}, expected {expected}"
                              + (f" at {decimals} decimals" if decimals is not None else ""))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Inputs:
    """Hash inputs under neutral aliases; do not export paths or workbook metadata."""
    def __init__(self):
        self.files = {}

    def record(self, alias, path):
        if not path.is_file():
            raise ValueError(f"Missing input: {alias}")
        self.files[alias] = (path, sha(path))
        return path

    def csv(self, alias, path):
        with self.record(alias, path).open(newline="", encoding="utf-8-sig") as stream:
            return list(csv.DictReader(stream))

    def verify_unchanged(self):
        for alias, (path, digest) in self.files.items():
            if sha(path) != digest:
                raise ValueError(f"Input changed during analysis: {alias}")

    def rows(self):
        return [dict(input=alias, sha256=digest) for alias, (_, digest) in self.files.items()]


def unique(rows, key, label):
    result = {}
    for row in rows:
        value = row[key]
        if value in result:
            raise ValueError(f"Duplicate {key} in {label}: {value}")
        result[value] = row
    return result


def number(value):
    if value in (None, "", "NJ", "NA"):
        return None
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("Non-finite numeric input")
    return result


def boolean(value):
    if value in (True, "True", "true", "1"):
        return True
    if value in (False, "False", "false", "0"):
        return False
    if value in (None, ""):
        return None
    raise ValueError("Unexpected boolean value")


def rating(value):
    if value == "NJ":
        return None
    result = number(value)
    if result not in (1, 2, 3, 4, 5):
        raise ValueError("Ratings must be 1..5 or explicit NJ; a blank is not NJ")
    return int(result)


def normalize_citation(value, rater):
    text = str(value or "").strip()
    if rater == "Expert B":
        text = B_CITATION_MAP.get(text, text)
    if text not in CITATIONS:
        raise ValueError(f"Missing or unsupported citation judgement for {rater}")
    return text


def read_forms(root, selection, key, inputs):
    forms, citations = {}, {item: dict(item_id=item, code=r["code"], corpus=r["corpus"],
                                     stratum=r["stratum"]) for item, r in key.items()}
    fixed_citations = None
    for index, rater in enumerate(RATERS, 1):
        path = inputs.record(f"{rater} workbook", root / f"results/expert-study/Expert-{index}.xlsx")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = {}
            for cells in workbook["1. Rate the models"].iter_rows(min_row=2, values_only=True):
                code = str(cells[0] or "").strip()
                if not code:
                    continue
                if code not in selection or code in values:
                    raise ValueError(f"Unknown or repeated rating code for {rater}")
                if str(cells[1]).strip().lower() != selection[code]["corpus"]:
                    raise ValueError(f"Rating corpus mismatch for {rater}: {code}")
                row = {q: rating(cells[2 + j]) for j, q in enumerate(DIMENSIONS)}
                decision = str(cells[7] or "").strip()
                if decision not in DECISIONS:
                    raise ValueError(f"Unsupported decision for {rater}: {code}")
                row.update(decision=decision, decision_score=DECISIONS[decision])
                values[code] = row
            if set(values) != set(selection):
                raise ValueError(f"Incomplete model ratings for {rater}")
            seen_ranks = set()
            for cells in workbook["2. Rank within domain"].iter_rows(values_only=True):
                code = str(cells[2] or "").strip() if len(cells) > 2 else ""
                if code not in selection:
                    continue
                rank = number(cells[3])
                if code in seen_ranks or rank not in (1, 2, 3, 4, 5):
                    raise ValueError(f"Invalid or duplicate rank for {rater}: {code}")
                values[code]["rank"] = int(rank)
                seen_ranks.add(code)
            if seen_ranks != set(selection):
                raise ValueError(f"Incomplete rankings for {rater}")
            seen_items, fixed = set(), {}
            for cells in workbook["3. Check citations"].iter_rows(min_row=2, values_only=True):
                if cells[0] is None:
                    continue
                raw_id = str(cells[0]).strip()
                item = raw_id if raw_id.startswith("C") else f"C{int(raw_id):02d}"
                if item not in key or item in seen_items:
                    raise ValueError(f"Unknown or repeated citation item for {rater}")
                if str(cells[1]).strip() != key[item]["code"]:
                    raise ValueError(f"Citation model mismatch: {item}")
                if str(cells[2]).strip().lower() != key[item]["corpus"]:
                    raise ValueError(f"Citation corpus mismatch: {item}")
                normalize_path = lambda x: " / ".join(p.strip() for p in str(x).replace("→", "/").split("/"))
                if normalize_path(cells[3]) != normalize_path(key[item]["feature"]):
                    raise ValueError(f"Citation feature mismatch: {item}")
                fixed[item] = tuple(cells[:8])
                label = normalize_citation(cells[8], rater)
                citations[item][rater + " label"] = label
                citations[item][rater + " score"] = CITATIONS[label]
                seen_items.add(item)
            if seen_items != set(key):
                raise ValueError(f"Incomplete citation sheet for {rater}")
            if fixed_citations is not None and fixed != fixed_citations:
                raise ValueError("The experts received different citation items")
            fixed_citations = fixed
            forms[rater] = values
        finally:
            workbook.close()
    return forms, list(citations.values())


def read_public_responses(directory, selection, key, inputs):
    """Use anonymised response CSVs, never private workbooks, for replication."""
    forms = {rater: {} for rater in RATERS}
    for row in inputs.csv("anonymised ratings", directory / "ratings.csv"):
        rater, code = row["rater"], row["code"]
        if rater not in forms or code not in selection or code in forms[rater]:
            raise ValueError("Unknown or duplicate anonymised rating identity")
        if row["corpus"] != selection[code]["corpus"]:
            raise ValueError("Anonymised rating corpus mismatch")
        values = {q: rating("NJ" if row[q] == "" else row[q]) for q in DIMENSIONS}
        decision = row["decision"]
        if decision not in DECISIONS or number(row["decision_score"]) != DECISIONS[decision]:
            raise ValueError("Anonymised decision coding mismatch")
        rank = number(row["rank"])
        if rank not in (1, 2, 3, 4, 5):
            raise ValueError("Invalid anonymised rank")
        values.update(decision=decision, decision_score=DECISIONS[decision], rank=int(rank))
        forms[rater][code] = values
    if any(set(rows) != set(selection) for rows in forms.values()):
        raise ValueError("Incomplete anonymised ratings")
    citations = inputs.csv("anonymised citations", directory / "citation_judgements.csv")
    indexed = unique(citations, "item_id", "anonymised citations")
    if set(indexed) != set(key):
        raise ValueError("Incomplete anonymised citations")
    for item, row in indexed.items():
        for field in ("code", "corpus", "stratum"):
            if row[field] != key[item][field]:
                raise ValueError("Anonymised citation identity mismatch")
        for rater in RATERS:
            label = row[rater + " label"]
            if label not in CITATIONS or number(row[rater + " score"]) != CITATIONS[label]:
                raise ValueError("Anonymised citation coding mismatch")
            row[rater + " score"] = CITATIONS[label]
    return forms, citations


def ordinal_alpha(a, b, domain):
    data = np.array([[np.nan if x is None else x for x in values] for values in (a, b)], dtype=float)
    return float(krippendorff.alpha(reliability_data=data, level_of_measurement="ordinal",
                                  value_domain=list(domain)))


def kendalls_w(rankings):
    data = np.array([rankdata(r, method="average") for r in rankings])
    m, n = data.shape
    ties = sum(sum(t ** 3 - t for t in Counter(r).values()) for r in data)
    denominator = m * m * (n ** 3 - n) - m * ties
    if denominator == 0:
        return None
    sums = data.sum(axis=0)
    return float(12 * ((sums - m * (n + 1) / 2) ** 2).sum() / denominator)


def tau(x, y):
    pairs = [(a, b) for a, b in zip(x, y) if a is not None and b is not None]
    if len(pairs) < 2:
        return dict(n=len(pairs), tau_b=None)
    a, b = zip(*pairs)
    value = float(kendalltau(a, b, variant="b").statistic)
    return dict(n=len(pairs), tau_b=value if math.isfinite(value) else None)


def rating_analyses(forms, selection, citations, checks):
    codes = list(selection)
    summaries, comparisons, alphas, decisions, concordance = [], [], [], [], []
    for rater in RATERS:
        for corpus in ("all", *CORPORA):
            subset = [c for c in codes if corpus == "all" or selection[c]["corpus"] == corpus]
            for dimension in ("pooled", *DIMENSIONS):
                qs = DIMENSIONS if dimension == "pooled" else (dimension,)
                values = [forms[rater][c][q] for c in subset for q in qs]
                present = [v for v in values if v is not None]
                summaries.append(dict(rater=rater, corpus=corpus, dimension=dimension,
                                      n=len(present), missing=len(values)-len(present),
                                      mean=statistics.fmean(present) if present else None))
    for rater, expected in zip(RATERS, (3.06, 1.72)):
        r = next(r for r in summaries if (r["rater"], r["corpus"], r["dimension"]) == (rater, "all", "pooled"))
        checks.equal(f"{rater} mean rating", r["mean"], expected, 2)
    for corpus in ("all", *CORPORA):
        cells = [(forms[RATERS[0]][c][q], forms[RATERS[1]][c][q]) for c in codes
                 if corpus == "all" or selection[c]["corpus"] == corpus for q in DIMENSIONS]
        paired = [(a, b) for a, b in cells if a is not None and b is not None]
        comparisons.append(dict(corpus=corpus, jointly_rated=len(paired),
                                B_lower=sum(b<a for a,b in paired), B_higher=sum(b>a for a,b in paired),
                                equal=sum(b==a for a,b in paired)))
    for field, expected in (("jointly_rated",46),("B_lower",34),("B_higher",1)):
        checks.equal(field, comparisons[0][field], expected)
    missing_b = [q for c in codes for q in DIMENSIONS if forms["Expert B"][c][q] is None]
    checks.equal("Expert B missing ratings", len(missing_b), 4)
    checks.equal("Expert B missing Q5 ratings", missing_b.count("Q5"), 4)
    expected_alpha = dict(Q1=-.335,Q2=-.116,Q3=-.314,Q4=-.059,Q5=-.457,pooled=-.098,
                          decisions=.163,citations=.196)
    for dimension in (*DIMENSIONS, "pooled", "decisions", "citations"):
        if dimension == "citations":
            matrix = [[r[rater+" score"] for r in citations] for rater in RATERS]
            domain = range(3)
        elif dimension == "decisions":
            matrix = [[forms[rater][c]["decision_score"] for c in codes] for rater in RATERS]
            domain = range(3)
        else:
            qs = DIMENSIONS if dimension == "pooled" else (dimension,)
            matrix = [[forms[rater][c][q] for c in codes for q in qs] for rater in RATERS]
            domain = range(1,6)
        value = ordinal_alpha(*matrix, domain)
        alphas.append(dict(dimension=dimension, units=len(matrix[0]),
                           jointly_rated=sum(a is not None and b is not None for a,b in zip(*matrix)),
                           ordinal_alpha=value, value_domain=":".join(map(str,domain))))
        checks.equal("ordinal alpha " + dimension, value, expected_alpha[dimension], 3)
    for corpus in CORPORA:
        subset = [c for c in codes if selection[c]["corpus"] == corpus]
        a,b = [[forms[rater][c]["decision_score"] for c in subset] for rater in RATERS]
        pairs = [(x,y) for x,y in zip(a,b) if x is not None and y is not None]
        agree = sum(x == y for x,y in pairs)
        decisions.append(dict(corpus=corpus, jointly_decided=len(pairs), agree=agree,
                              agreement=agree/len(pairs) if pairs else None))
        checks.equal(corpus + " decision agreement", agree, 4 if corpus == "repair" else 1)
        rankings = [[forms[rater][c]["rank"] for c in subset] for rater in RATERS]
        w = kendalls_w(rankings)
        concordance.append(dict(corpus=corpus,n=len(subset),raters=2,kendalls_W=w,
                                inter_rater_tau_b=tau(*rankings)["tau_b"]))
        checks.equal(corpus+" Kendall W", w, .855 if corpus == "repair" else .526, 3)
        if corpus == "federation":
            checks.equal("Federation inter-rater tau_b", concordance[-1]["inter_rater_tau_b"], 0., 3)
    return dict(rating_summary=summaries, rating_comparisons=comparisons, ordinal_alpha=alphas,
                decision_agreement=decisions, ranking_concordance=concordance)


def xml_shape(path):
    root = ET.parse(path).getroot()
    struct = root.find("struct")
    if struct is None:
        raise ValueError("Selected XML has no struct element")
    nodes = [n for n in struct.iter() if n.tag in FEATURE_TAGS]
    if not nodes:
        raise ValueError("Selected XML has no feature tree")
    internal = [n for n in nodes if any(c.tag in FEATURE_TAGS for c in n)]
    mandatory = sum(n.get("mandatory", "false").lower() == "true" for n in nodes[1:])
    return dict(size=len(nodes), internal_features=len(internal),
                internal_and=sum(n.tag=="and" for n in internal),
                internal_or=sum(n.tag=="or" for n in internal),
                internal_alt=sum(n.tag=="alt" for n in internal), mandatory_children=mandatory,
                empty_groups_ignored=sum(n.tag in {"and","or","alt"} and n not in internal for n in nodes),
                degeneracy=bool(internal) and all(n.tag=="and" for n in internal) and mandatory==0)


def selected_xml(root, row):
    # Frozen manifests contain historical absolute paths. Restore under this checkout.
    value = row["xml_path"].replace("\\", "/")
    if "/results/" in value:
        value = "results/" + value.split("/results/",1)[1]
    path = (root / value).resolve()
    if not path.is_relative_to(root.resolve() / "results"):
        raise ValueError("Selected XML must resolve inside results")
    return path


def automated_scores(root, selection, forms, inputs):
    base = root / "results/ifs-2027/analysis/reproduction-2026-10-07"
    wide = unique(inputs.csv("campaign aggregate",base/"aggregate-current-v4/wide.csv"),"run_id","wide")
    matching = [r for r in inputs.csv("campaign matching",base/"matching-comparison-v2/matching_comparison.csv") if float(r["tau"])==.4]
    matching = unique(matching,"run_id","matching at tau 0.4")
    siblings = {}
    for row in inputs.csv("campaign sibling agreement",base/"sibling-current-v1/sibling_metrics.csv"):
        if row["matching_policy"] not in POLICIES or float(row["tau"])!=.4:
            continue
        key = (row["run_id"],row["matching_policy"])
        if key in siblings:
            raise ValueError("Duplicate sibling score")
        siblings[key] = row
    astra = unique(inputs.csv("Astra N1 evaluation",root/"results/expert-study/analysis/astra-n1-rag-evaluation-v1.csv"),"run_id","Astra extension")
    merged, shapes = [], []
    for code, selected in selection.items():
        rid = selected["run_id"]
        extension = selected["model_id"] == "gpt-6-astra"
        meta = astra[rid] if extension else wide[rid]
        score = meta if extension else matching[rid]
        for field in ("model_id","corpus","N","seed"):
            if str(meta[field])!=str(selected[field]):
                raise ValueError(f"Selected metadata mismatch for {code}: {field}")
        if not boolean(meta["metamodel_block"]) or meta["grounding"]!="rag" or meta["ordering_id"]!="primary" or int(meta["k_doc"])!=5:
            raise ValueError(f"Selected output is outside the guided-RAG cell: {code}")
        if (extension and not boolean(meta["completed"])) or (not extension and meta["inventory_status"]!="completed"):
            raise ValueError(f"Selected output is not completed: {code}")
        if extension and float(meta["tau"])!=.4:
            raise ValueError("Astra extension threshold differs from 0.4")
        path = inputs.record("XML " + code, selected_xml(root,selected))
        if sha(path)!=selected["xml_sha256"]:
            raise ValueError(f"Frozen XML hash mismatch: {code}")
        shape = xml_shape(path)
        shapes.append(dict(code=code,corpus=selected["corpus"],**shape))
        size = int(meta["semantic__n_generated"])
        if shape["size"]!=size:
            raise ValueError(f"XML and scored feature counts differ: {code}")
        row = dict(code=code,corpus=selected["corpus"],model=MODELS[selected["model_id"]],
                   model_id=selected["model_id"],N=int(selected["N"]),seed=int(selected["seed"]),run_id=rid,
                   size=size,conformance=boolean(meta["structural__structural_conformance"]),
                   satisfiability=boolean(meta["structural__satisfiable"]),
                   satisfiability_status=meta["structural__satisfiable__status"],degeneracy=shape["degeneracy"])
        for policy in POLICIES:
            for suffix in ("precision","recall_total","f1_total"):
                row[policy+"__"+suffix] = number(score[policy+"__"+suffix])
            skey = "siblings_"+policy+"__sibling_precision"
            row[policy+"__sibling_precision"] = number(meta[skey] if extension else siblings[(rid,policy)]["sibling_precision"])
        for field,suffix in (("semantic__semantic_precision","precision"),("semantic__semantic_recall_total","recall_total"),("semantic__semantic_f1_total","f1_total")):
            if not math.isclose(row["independent_max__"+suffix],float(meta[field]),abs_tol=1e-12):
                raise ValueError(f"Matching and aggregate scores differ: {code}")
        denominator = row["one_to_one__precision"]
        row["precision_ratio"] = row["independent_max__precision"]/denominator if denominator else None
        for rater in RATERS:
            row.update({rater+" "+k:v for k,v in forms[rater][code].items()})
        row["mean_expert_rank"] = statistics.fmean(row[r+" rank"] for r in RATERS)
        merged.append(row)
    return merged, shapes, wide, matching, astra


def correlations(merged, checks):
    ranks, ratings = [], []
    for corpus in CORPORA:
        rows = [r for r in merged if r["corpus"]==corpus]
        for rater in (*RATERS,"Mean rank"):
            key = "mean_expert_rank" if rater=="Mean rank" else rater+" rank"
            for policy in POLICIES:
                ranks.append(dict(corpus=corpus,rater=rater,policy=policy,
                                  **tau([-r[key] for r in rows],[r[policy+"__f1_total"] for r in rows])))
        for rater in RATERS:
            for policy in POLICIES:
                for dimension,metric in (("Q1","precision"),("Q2","recall_total"),("Q3","sibling_precision"),("Q5","minus_precision_ratio")):
                    y = [-r["precision_ratio"] if r["precision_ratio"] is not None else None for r in rows] if dimension=="Q5" else [r[policy+"__"+metric] for r in rows]
                    ratings.append(dict(corpus=corpus,rater=rater,dimension=dimension,policy=policy,
                                        measure=metric,**tau([r[rater+" "+dimension] for r in rows],y)))
    expected = {("repair","independent_max"):(.53,.74,.74),("repair","one_to_one"):(.32,.53,.53),
                ("federation","independent_max"):(-.20,-.22,0.),("federation","one_to_one"):(.20,.22,.40)}
    for (corpus,policy), values in expected.items():
        for rater,value in zip((*RATERS,"Mean rank"),values):
            r=next(r for r in ranks if (r["corpus"],r["policy"],r["rater"])==(corpus,policy,rater))
            checks.equal(f"rank tau {corpus} {policy} {rater}",r["tau_b"],value,2)
    for corpus,policy,value in (("repair","independent_max",.77),("repair","one_to_one",.82),
                                 ("federation","independent_max",.89),("federation","one_to_one",.89)):
        r=next(r for r in ratings if (r["corpus"],r["policy"],r["rater"],r["dimension"])==(corpus,policy,"Expert A","Q2"))
        checks.equal(f"Q2 recall {corpus} {policy}",r["tau_b"],value,2)
    q1=[r["tau_b"] for r in ratings if r["dimension"]=="Q1"]
    checks.equal("Q1 precision minimum",min(q1),-.60,2)
    checks.equal("Q1 precision maximum",max(q1),.52,2)
    return ranks,ratings


def discordances(merged,checks):
    result=[]
    for corpus in CORPORA:
        rows=sorted((r for r in merged if r["corpus"]==corpus),key=lambda r:r["code"])
        for a,b in combinations(rows,2):
            d1=a["independent_max__f1_total"]-b["independent_max__f1_total"]
            d2=a["one_to_one__f1_total"]-b["one_to_one__f1_total"]
            if d1*d2 >= 0:
                continue
            row=dict(corpus=corpus,code_1=a["code"],code_2=b["code"],
                     independent_max_preferred=a["code"] if d1>0 else b["code"],
                     one_to_one_preferred=a["code"] if d2>0 else b["code"])
            for rater in RATERS:
                diff=b[rater+" rank"]-a[rater+" rank"]
                row[rater+" agrees_with"]="tie" if diff==0 else ("independent_max" if diff*d1>0 else "one_to_one")
            result.append(row)
    checks.equal("discordant pairs",len(result),5)
    expected={frozenset(p):"one_to_one" for p in [("Rep-28","Rep-79"),("Fed-97","Fed-44"),("Fed-97","Fed-67")]}
    expected.update({frozenset(p):"independent_max" for p in [("Rep-28","Rep-57"),("Rep-79","Rep-57")]})
    checks.equal("discordant pair identities",{frozenset((r["code_1"],r["code_2"])) for r in result}==set(expected),True)
    for r in result:
        for rater in RATERS:
            checks.equal(f"discordance {r['code_1']}/{r['code_2']} {rater}",r[rater+" agrees_with"],expected[frozenset((r["code_1"],r["code_2"]))])
    return result


def variability(merged,checks):
    result=[]
    for rater,values in zip(RATERS,((1.67,2.57),(1.,1.57))):
        for degenerate,expected in zip((True,False),values):
            group=[r for r in merged if r["degeneracy"]==degenerate]
            scores=[r[rater+" Q4"] for r in group if r[rater+" Q4"] is not None]
            value=statistics.fmean(scores)
            result.append(dict(rater=rater,degenerate=degenerate,n_outputs=len(group),n_ratings=len(scores),mean_Q4=value))
            checks.equal(f"Q4 mean {rater} degenerate={degenerate}",value,expected,2)
    return result


def citation_analyses(rows,checks):
    result=[]
    for rater in RATERS:
        for corpus in ("all",*CORPORA):
            for stratum in ("agree","disagree"):
                rs=[r for r in rows if r["stratum"]==stratum and (corpus=="all" or r["corpus"]==corpus)]
                counts=Counter(r[rater+" label"] for r in rs)
                judged=[r[rater+" score"] for r in rs if r[rater+" score"] is not None]
                supported=sum(v>=1 for v in judged)
                result.append(dict(rater=rater,corpus=corpus,stratum=stratum,items=len(rs),decided=len(judged),
                                   cannot_judge=counts["Cannot judge"],yes=counts["Yes"],partially=counts["Partially"],no=counts["No"],
                                   at_least_partial=supported,support_rate=supported/len(judged) if judged else None))
    for rater,stratum,n,supported in (("Expert A","agree",6,6),("Expert A","disagree",6,5),
                                      ("Expert B","agree",5,4),("Expert B","disagree",5,4)):
        r=next(r for r in result if (r["rater"],r["corpus"],r["stratum"])==(rater,"all",stratum))
        checks.equal(f"citation denominator {rater} {stratum}",r["decided"],n)
        checks.equal(f"citation support {rater} {stratum}",r["at_least_partial"],supported)
        checks.equal(f"citation cannot judge {rater} {stratum}",r["cannot_judge"],0 if rater=="Expert A" else 1)
    jointly=[r for r in rows if all(r[t+" score"] is not None for t in RATERS)]
    agree=sum(r["Expert A score"]==r["Expert B score"] for r in jointly)
    checks.equal("citations jointly decided",len(jointly),10)
    checks.equal("citations exact agreement",agree,3)
    return result,[dict(items=len(rows),jointly_decided=len(jointly),exact_agreement=agree,agreement=agree/len(jointly))]


def mid_percentile(value, population):
    """Include the selected observation itself; comparisons use saved full precision."""
    if not population or value not in population:
        raise ValueError("Percentile population must contain the selected value")
    return 100*(sum(v<value for v in population)+.5*sum(v==value for v in population))/len(population)


def percentiles(merged,wide,matching,astra,checks):
    result=[]
    for row in merged:
        extension=row["model"]=="Astra"
        candidates=astra.values() if extension else wide.values()
        peers=[]
        for r in candidates:
            if (r["model_id"],r["corpus"],int(r["N"]))!=(row["model_id"],row["corpus"],row["N"]):continue
            if r["grounding"]!="rag" or not boolean(r["metamodel_block"]) or r["ordering_id"]!="primary" or int(r["k_doc"])!=5:continue
            if extension:
                if not boolean(r["completed"]) or float(r["tau"])!=.4:continue
            elif not r["arm"].startswith("guided_") or r["inventory_status"]!="completed":continue
            if not extension and r["run_id"] not in matching:
                raise ValueError("A completed percentile peer has no matching scores")
            peers.append(r)
        if row["run_id"] not in {r["run_id"] for r in peers}:
            raise ValueError("Selected output absent from percentile cell")
        values={"size":[int(r["semantic__n_generated"]) for r in peers]}
        for policy in POLICIES:
            values[policy+"__f1_total"]=[float((r if extension else matching[r["run_id"]])[policy+"__f1_total"]) for r in peers]
        record={k:row[k] for k in ("code","corpus","model","N","run_id","size")}
        record["cell_n"]=len(peers)
        for metric,vs in values.items():
            record[metric+"_percentile"]=mid_percentile(row[metric],vs)
        result.append(record)
    for corpus,expected in (("repair",(70,70,10)),("federation",(90,90,70))):
        row=next(r for r in result if r["model"]=="Astra" and r["corpus"]==corpus)
        for field,value in zip(("size_percentile","independent_max__f1_total_percentile","one_to_one__f1_total_percentile"),expected):
            checks.equal("Table 12 Astra "+corpus+" "+field,row[field],value,0)
    return result


def fmt(value):
    if value is None:return ""
    if isinstance(value,(float,np.floating)):return f"{value:.3f}"
    return str(value)


def write_csv(path,rows):
    if not rows:raise ValueError("Refusing to write an empty analysis table")
    with path.open("x",newline="",encoding="utf-8") as stream:
        writer=csv.DictWriter(stream,fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows({k:fmt(v) for k,v in row.items()} for row in rows)


def md_table(rows,fields=None):
    fields=fields or list(rows[0])
    return "\n".join(["| "+" | ".join(fields)+" |","| "+" | ".join("---" for _ in fields)+" |"]+
                      ["| "+" | ".join(fmt(r[k]) if r[k] is not None else "NA" for k in fields)+" |" for r in rows])


def summary(tables,checks):
    sections=["# Expert assessment — final analysis", "All supplied numerical sanity checks passed. Inputs were read-only and hash-verified. "
              "Only Expert A and Expert B identify raters; free-text comments and workbook metadata are not exported.",
              "## Section 4.5: population and definitions", "Ten selected guided-RAG outputs, five per corpus, rated by two experts. "
              "The Astra N=1 extension remains separate from the main campaign. Ratings NJ and citations Cannot judge are missing, not zero. "
              "Expert B labels are mapped exactly as authorised: OK → Yes; Not exactly → Partially; "
              "not sure, depends on the context → Cannot judge; No → No. Decisions use Reject=0, Accept with revision=1, Accept=2.",
              "Krippendorff alpha uses ordinal distances, value_domain 1..5 for Q1–Q5 and pooled ratings, "
              "and 0..2 for decisions/citations. Pooled ratings flatten the five dimensions within each output. "
              "Kendall W uses within-rater average ranks and tie correction. Kendall tau-b is pairwise complete; "
              "rank–F1 correlations negate ranks so larger values consistently mean better. Mean rank is the arithmetic "
              "mean of the two submitted ranks. Q5 is correlated with minus (independent-max precision / one-to-one precision); "
              "this predictor is shared, so its two policy rows are identical by definition.",
              "## Section 5.5: ratings and agreement",md_table(tables["rating_summary"]),md_table(tables["rating_comparisons"]),
              md_table(tables["ordinal_alpha"]),md_table(tables["decision_agreement"]),md_table(tables["ranking_concordance"]),
              "## Rankings versus automated scores",md_table(tables["ranking_correlations"]),
              "## Ratings versus automated measures",md_table(tables["rating_metric_correlations"]),
              "## Policy-discordant pairs",md_table(tables["discordant_pairs"]),
              "## Table 4: selected outputs",md_table(tables["table4" ]),
              "Degeneracy is computed from the unchanged XML: every internal feature is AND and no non-root feature is mandatory. "
              "Empty group leaves are ignored only when identifying internal features. This supplies Rep-28's degeneracy "
              "without treating its nonconformance as conformance. Unavailable satisfiability remains blank/NA, not false.",
              md_table(tables["variability_summary"]),
              "## Citation assessment",md_table(tables["citation_summary"]),md_table(tables["citation_agreement"]),
              "Citation strata denote agreement with reference attribution, not human ground truth. Support rates exclude Cannot judge. "
              "The 12 selected items are a stratified audit, not a representative estimate over all generated citations.",
              "## Appendix E, Table 12: within-cell percentiles",md_table(tables["table12_percentiles"]),
              "Percentiles use 100 × (number lower + 0.5 × number equal) / n, including the selected run itself. "
              "Cells contain completed scored outputs of the same corpus, model, N, guided-RAG condition, primary ordering and k_doc=5; "
              "there is no conformance filter. Comparisons use full-precision saved scores, not rounded displays. "
              "Only the two supplied Astra percentile triples were available as published-value checks; the other rows are computed, "
              "not certified against an unseen manuscript table.",
              "## Interpretation limits", "These are descriptive results from two raters and ten selected outputs. "
              "No population-wide model ranking, causal policy advantage or statistical significance is established. "
              "No bootstrap confidence intervals were requested for this small expert sample. "
              "All reported calculations are deterministic; no LLM or embedding calls are made.",
              "## Reproducibility",f"Passed checks: {len(checks.rows)}. See sanity_checks.csv, input_hashes.csv and software_versions.csv. "
              "CSV scores are displayed to three decimals; computations use unrounded values. "
              "expert_assessment_merged.csv has one row per selected output; ratings.csv and citation_judgements.csv retain the normalised response-level data."]
    return "\n\n".join(sections)+"\n"


def analyse(root,inputs,checks,public_inputs=None):
    base=root/"results/expert-study"
    selection_path = public_inputs/"model_selection.csv" if public_inputs else base/"selection-v3/expert_sample_frozen.csv"
    selection=unique(inputs.csv("frozen selection",selection_path),"code","selection")
    checks.equal("selected models",len(selection),10)
    for corpus in CORPORA:checks.equal(corpus+" selected models",sum(r["corpus"]==corpus for r in selection.values()),5)
    unique(list(selection.values()),"run_id","selection")
    key_path = public_inputs/"citation_items.csv" if public_inputs else base/"citations-v3/citation_key_AUTHOR_ONLY.csv"
    key=unique(inputs.csv("citation strata",key_path),"item_id","citation key")
    checks.equal("citation items",len(key),12)
    for r in key.values():
        if r["code"] not in selection or r["run_id"]!=selection[r["code"]]["run_id"]:
            raise ValueError("Citation selection/run identity mismatch")
    for corpus in CORPORA:
        for stratum in ("agree","disagree"):
            checks.equal(corpus+" citation stratum "+stratum,sum((r["corpus"],r["stratum"])==(corpus,stratum) for r in key.values()),3)
    forms,citations=(read_public_responses(public_inputs,selection,key,inputs) if public_inputs
                     else read_forms(root,selection,key,inputs))
    tables=rating_analyses(forms,selection,citations,checks)
    merged,shapes,wide,matching,astra=automated_scores(root,selection,forms,inputs)
    tables["ranking_correlations"],tables["rating_metric_correlations"]=correlations(merged,checks)
    tables["discordant_pairs"]=discordances(merged,checks)
    tables["variability_summary"]=variability(merged,checks)
    tables["citation_summary"],tables["citation_agreement"]=citation_analyses(citations,checks)
    tables["table12_percentiles"]=percentiles(merged,wide,matching,astra,checks)
    tables["expert_assessment_merged"]=merged
    tables["degeneracy_xml_audit"]=shapes
    tables["citation_judgements"]=citations
    tables["ratings"]=[dict(rater=rater,code=c,corpus=selection[c]["corpus"],**forms[rater][c]) for rater in RATERS for c in selection]
    columns=["code","corpus","model","N","size","conformance","satisfiability","degeneracy"]
    columns += [p+"__f1_total" for p in POLICIES]
    columns += [r+" "+q for r in RATERS for q in (*DIMENSIONS,"decision","rank")]
    tables["table4"]=[{k:r[k] for k in columns} for r in merged]
    inputs.verify_unchanged()
    return tables


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root",type=Path,default=ROOT,help="Checkout containing the unchanged inputs")
    parser.add_argument("--output",type=Path,default=None,help="New directory under results/expert-study")
    parser.add_argument("--public-inputs",type=Path,help="Directory containing anonymised ratings, citations and selection CSVs; no workbooks read")
    args=parser.parse_args(argv)
    root=args.root.resolve()
    out=(args.output or root/"results/expert-study/analysis-final").resolve()
    allowed=root/"results/expert-study"
    if out.exists() or out==allowed or not out.is_relative_to(allowed):
        parser.error("Choose a NEW output directory inside results/expert-study; existing outputs are never overwritten")
    inputs,checks=Inputs(),Checks()
    try:
        tables=(analyse(root,inputs,checks,args.public_inputs.resolve()) if args.public_inputs
                else analyse(root,inputs,checks))
    except (ValueError,KeyError,OSError,ET.ParseError) as exc:
        # Do not echo exception filenames: they can contain personal path components.
        detail=str(exc) if isinstance(exc,ValueError) else "Required input could not be read or did not have the expected schema"
        out.mkdir(parents=True)
        if checks.rows:write_csv(out/"sanity_checks.csv",checks.rows)
        if inputs.files:write_csv(out/"input_hashes.csv",inputs.rows())
        (out/"summary.md").write_text("# Expert assessment — STOPPED\n\n"+detail+"\n\nNo final analysis tables were written. No inputs were modified.\n",encoding="utf-8")
        print("STOPPED: "+detail,file=sys.stderr)
        return 1
    tables["sanity_checks"]=checks.rows
    tables["input_hashes"]=inputs.rows()
    tables["software_versions"]=[dict(component=p,version=version(p)) for p in ("numpy","scipy","openpyxl","krippendorff")]
    tables["software_versions"].append(dict(component="expert_analysis.py SHA256",version=sha(Path(__file__))))
    out.mkdir(parents=True)
    for name,rows in tables.items():write_csv(out/(name+".csv"),rows)
    (out/"summary.md").write_text(summary(tables,checks),encoding="utf-8")
    print(f"Passed {len(checks.rows)} sanity checks; wrote {len(tables)} CSVs and summary.md.")
    return 0


if __name__=="__main__":
    raise SystemExit(main())
