#!/usr/bin/env python3
"""Protocol-v3 CLI; legacy v1 helpers retained for historical regression tests."""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import random
import re
import zipfile
from collections import defaultdict
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook


REPO = Path(__file__).resolve().parents[1]
MODELS = {
    "deepseek-v4.1-flash:cloud": "deepseek_flash",
    "glm-5.3-flash:cloud": "glm_flash",
    "gpt-6-astra": "astra",
}
CORPORA = ("repair", "federation")
MATCH_FIELDS = (
    "N", "grounding", "metamodel_block", "k_doc", "ordering_id",
    "max_output_tokens", "prompt_template_hash", "metamodel_hash",
    "chunks_hash", "encoder_digest", "root_feature",
)
TRACE = re.compile(r"\s*Trace:\s*\[[^\]]*\]\s*$", re.IGNORECASE)
ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"


def _parse_xml(path: Path):
    return etree.parse(str(path), etree.XMLParser(resolve_entities=False, no_network=True))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def _candidate(meta_path: Path, expected_run_id: str | None = None) -> dict:
    meta = _read_json(meta_path)
    cfg = meta["config"]
    if expected_run_id and meta.get("run_id") != expected_run_id:
        raise ValueError(f"inventory/run_meta run_id mismatch: {meta_path}")
    xml = meta_path.parent / "fm_gen.xml"
    if not meta.get("completed") or meta.get("terminal_status") != "completed" or not xml.is_file():
        raise ValueError(f"run is not complete with a final XML: {meta_path}")
    root = _parse_xml(xml).getroot()
    if etree.QName(root).localname != "featureModel" or len(root.findall("struct")) != 1:
        raise ValueError(f"not an extractable FeatureIDE tree: {xml}")
    struct = root.find("struct")
    if struct is None or len([n for n in struct if isinstance(n.tag, str)]) != 1:
        raise ValueError(f"not exactly one structural root: {xml}")
    if (cfg["N"], cfg["grounding"], cfg["metamodel_block"]) != (1, "rag", True) or cfg["model_id"] not in MODELS:
        raise ValueError(f"wrong expert-study configuration: {meta_path}")
    return {
        "model_key": MODELS[cfg["model_id"]],
        "model_id": cfg["model_id"],
        "corpus": cfg["corpus"],
        "seed": int(cfg["seed"]),
        "run_id": meta["run_id"],
        "xml_path": str(xml.resolve()),
        "xml_sha256": _sha256(xml),
        "run_meta_path": str(meta_path.resolve()),
        "config": cfg,
    }


def _load_candidates(open_inventory: Path, astra_results: Path) -> dict[tuple[str, str, int], dict]:
    candidates = {}
    for row in _read_json(open_inventory):
        if not (
            row.get("arm") == "guided_baseline"
            and row.get("N") == 1
            and row.get("grounding") == "rag"
            and row.get("metamodel_block") is True
            and row.get("model_id") in MODELS
            and row.get("model_id") != "gpt-6-astra"
            and row.get("corpus") in CORPORA
            and row.get("seed") in range(5)
            and row.get("completed") is True
        ):
            continue
        meta_path = REPO / row["resolved_path"] / "run_meta.json"
        item = _candidate(meta_path, row["run_id"])
        key = (item["model_key"], item["corpus"], item["seed"])
        if key in candidates:
            raise ValueError(f"duplicate open-model candidate: {key}")
        candidates[key] = item
    if not astra_results.is_dir():
        raise FileNotFoundError(f"Astra results directory is missing: {astra_results}")
    for meta_path in sorted(astra_results.glob("*/**/run_meta.json")):
        meta = _read_json(meta_path)
        cfg = meta.get("config", {})
        if not (
            cfg.get("model_id") == "gpt-6-astra"
            and cfg.get("N") == 1
            and cfg.get("grounding") == "rag"
            and cfg.get("metamodel_block") is True
            and cfg.get("corpus") in CORPORA
            and cfg.get("seed") in range(5)
            and meta.get("completed") is True
        ):
            continue
        item = _candidate(meta_path)
        key = (item["model_key"], item["corpus"], item["seed"])
        if key in candidates:
            raise ValueError(f"duplicate completed Astra candidate: {key}")
        candidates[key] = item
    return candidates


def _astra_attempts(astra_results: Path) -> list[dict]:
    attempts = []
    for meta_path in sorted(astra_results.glob("*/**/run_meta.json")):
        meta = _read_json(meta_path)
        cfg = meta.get("config", {})
        if not (
            cfg.get("model_id") == "gpt-6-astra" and cfg.get("N") == 1
            and cfg.get("grounding") == "rag" and cfg.get("metamodel_block") is True
            and cfg.get("corpus") in CORPORA and cfg.get("seed") in range(5)
        ):
            continue
        attempts.append({
            "corpus": cfg["corpus"], "seed": int(cfg["seed"]),
            "run_id": meta.get("run_id", ""),
            "completed": bool(meta.get("completed")),
            "terminal_status": meta.get("terminal_status", ""),
            "failed_step": meta.get("failed_step", ""),
            "run_meta_path": str(meta_path.resolve()),
        })
    keys = [(row["corpus"], row["seed"]) for row in attempts]
    expected = {(corpus, seed) for corpus in CORPORA for seed in range(5)}
    if len(keys) != len(set(keys)):
        raise ValueError("multiple Astra attempts share a corpus/seed; resolve and document them before selection")
    if set(keys) != expected:
        raise ValueError(f"Astra attempt records are incomplete; missing {sorted(expected - set(keys))}")
    return attempts


def _new_code(rng: random.Random, used: set[str], prefix: str) -> str:
    while True:
        code = prefix + "".join(rng.choice(ALPHABET) for _ in range(5))
        if code not in used:
            used.add(code)
            return code


def _doc_index() -> dict[str, dict]:
    result = {}
    for path in (
        REPO / "data/raw/repair/manifest_repair.csv",
        REPO / "data/raw/federation/manifest_fed.csv",
    ):
        with path.open(newline="", encoding="utf-8-sig") as handle:
            for row in csv.DictReader(handle, delimiter=";"):
                doc_id = row["doc_id"].strip()
                if doc_id in result:
                    raise ValueError(f"duplicate source document ID: {doc_id}")
                result[doc_id] = row
    return result


def _feature_citations(xml_path: Path) -> list[tuple[str, str]]:
    root = _parse_xml(xml_path).getroot()
    edges = []
    feature_tags = {"and", "or", "alt", "feature"}
    for node in root.xpath("./struct//*[self::and or self::or or self::alt or self::feature]"):
        name = node.get("name")
        desc = node.findtext("description") or ""
        marker = re.search(r"Trace:\s*\[([^\]]*)\]\s*$", desc, re.IGNORECASE)
        if not name or not marker:
            continue
        path_parts = []
        ancestor = node
        while ancestor is not None and etree.QName(ancestor).localname in feature_tags:
            path_parts.append(ancestor.get("name", "(unnamed)"))
            ancestor = ancestor.getparent()
        feature_path = " / ".join(reversed(path_parts))
        for doc_id in re.findall(r"(?:rep|fed)_\d{2}", marker.group(1)):
            edges.append((feature_path, doc_id))
    return sorted(set(edges))


def _write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows([{key: row.get(key, "") for key in columns} for row in rows])


def select(args: argparse.Namespace) -> None:
    output = args.output.resolve()
    if output.exists():
        raise FileExistsError(f"refusing to overwrite selection: {output}")
    astra_attempts = _astra_attempts(args.astra_results)
    candidates = _load_candidates(args.open_inventory, args.astra_results)
    selected = []
    seed_policy = {}
    for corpus in CORPORA:
        complete = [
            seed for seed in range(5)
            if all((model, corpus, seed) in candidates for model in MODELS.values())
        ]
        if len(complete) < 4:
            raise ValueError(f"only {len(complete)} complete three-model triplets for {corpus}: {complete}")
        seeds = complete[:4]
        seed_policy[corpus] = {"eligible": complete, "selected": seeds, "reserve_used": 4 in seeds}
        for seed in seeds:
            triplet = [candidates[(model, corpus, seed)] for model in MODELS.values()]
            for field in MATCH_FIELDS:
                values = {json.dumps(item["config"].get(field), sort_keys=True) for item in triplet}
                if len(values) != 1:
                    raise ValueError(f"configuration mismatch for {corpus} seed {seed}: {field}={values}")
            selected.extend(triplet)
    rng = random.Random(args.blind_seed)
    used = set()
    set_codes = {}
    for item in selected:
        item["code"] = _new_code(rng, used, "FM-")
        key = (item["corpus"], item["seed"])
        if key not in set_codes:
            set_codes[key] = _new_code(rng, used, "SET-")
        item["set_code"] = set_codes[key]
    selected.sort(key=lambda item: (item["corpus"], item["seed"], item["model_key"]))
    index = _doc_index()
    citations = []
    for item in selected:
        for feature, doc_id in _feature_citations(Path(item["xml_path"])):
            if doc_id not in index:
                continue  # An L1-invalid marker is not a document-support judgment.
            citations.append({
                "run_id": item["run_id"], "corpus": item["corpus"],
                "model_key": item["model_key"], "feature": feature,
                "doc_id": doc_id, "doc_title": index[doc_id].get("title", ""),
                "excerpt": "", "excerpt_origin": "", "excerpt_locator": "",
            })
    selection = {
        "version": 2,
        "purpose": "blinded configured-pipeline comparison; not a Top-FM sample",
        "blind_seed": args.blind_seed,
        "open_inventory": str(args.open_inventory.resolve()),
        "open_inventory_sha256": _sha256(args.open_inventory),
        "astra_results": str(args.astra_results.resolve()),
        "astra_attempts": astra_attempts,
        "seed_policy": seed_policy,
        "items": selected,
    }
    # Balance the 15 document-support items across providers and domains, with
    # at most one sampled citation edge per FM. This is not a prevalence sample.
    citation_rng = random.Random(args.blind_seed + 17)
    quotas = {
        ("deepseek_flash", "repair"): 3, ("deepseek_flash", "federation"): 2,
        ("glm_flash", "repair"): 2, ("glm_flash", "federation"): 3,
        ("astra", "repair"): 3, ("astra", "federation"): 2,
    }
    sampled = []
    citation_issues = []
    for key, quota in quotas.items():
        pool = [row for row in citations if (row["model_key"], row["corpus"]) == key]
        by_run = defaultdict(list)
        for row in pool:
            by_run[row["run_id"]].append(row)
        if len(by_run) < quota:
            citation_issues.append({"model_key": key[0], "corpus": key[1], "available_fms": len(by_run), "required_fms": quota})
            continue
        # Sample FMs uniformly first, then one cited edge uniformly per FM.
        for run_id in citation_rng.sample(sorted(by_run), quota):
            sampled.append(citation_rng.choice(by_run[run_id]))
    citation_rng.shuffle(sampled)
    output.mkdir(parents=True)
    (output / "selection.json").write_text(json.dumps(selection, indent=2) + "\n", encoding="utf-8")
    _write_csv(
        output / "astra_attempts.csv",
        ["corpus", "seed", "run_id", "completed", "terminal_status", "failed_step", "run_meta_path"],
        astra_attempts,
    )
    _write_csv(
        output / "citation_candidates.csv",
        ["run_id", "corpus", "model_key", "feature", "doc_id", "doc_title", "excerpt", "excerpt_origin", "excerpt_locator"],
        citations,
    )
    if citation_issues:
        (output / "citation_sampling_issue.json").write_text(json.dumps(citation_issues, indent=2) + "\n", encoding="utf-8")
        print("citation audit cannot use the planned balanced 15-item design; amend its scope before packet generation")
    else:
        _write_csv(
            output / "citation_items_to_complete.csv",
            ["run_id", "corpus", "model_key", "feature", "doc_id", "doc_title", "excerpt", "excerpt_origin", "excerpt_locator"],
            sampled,
        )
    print(f"selected {len(selected)} FMs in 8 complete triplets; {len(citations)} real-document citation candidates")
    print(f"AUTHOR-ONLY: {output}")


def _constraint(node) -> str:
    tag = etree.QName(node).localname
    children = [child for child in node if isinstance(child.tag, str)]
    if tag == "var":
        return node.text or "?"
    if tag == "not" and len(children) == 1:
        return f"NOT {_constraint(children[0])}"
    ops = {"imp": " IMPLIES ", "eq": " IFF ", "conj": " AND ", "disj": " OR "}
    if tag in ops and children:
        return "(" + ops[tag].join(_constraint(child) for child in children) + ")"
    return etree.tostring(node, encoding="unicode", with_tail=False)


def _render_node(node) -> str:
    tag = etree.QName(node).localname
    label = html.escape(node.get("name", "(unnamed)"))
    mandatory = node.get("mandatory")
    qualifiers = [tag.upper()]
    if node.getparent() is not None and node.getparent().tag != 'struct':
        qualifiers.append("mandatory" if mandatory in ("true", "1") else "optional")
    desc = TRACE.sub("", node.findtext("description") or "").strip()
    detail = f'<div class="description">{html.escape(desc)}</div>' if desc else ""
    children = [child for child in node if isinstance(child.tag, str) and etree.QName(child).localname in ("and", "or", "alt", "feature")]
    nested = "<ul>" + "".join(_render_node(child) for child in children) + "</ul>" if children else ""
    heading=f'<strong>{label}</strong> <small>[{"; ".join(qualifiers)}]</small>'
    return (f'<li><details open><summary>{heading}</summary>{detail}{nested}</details></li>'
            if children else f'<li>{heading}{detail}</li>')


def _render_html(item: dict, display_code: str) -> str:
    root = _parse_xml(Path(item["xml_path"])).getroot()
    struct = root.find("struct")
    assert struct is not None
    feature_root = next(child for child in struct if isinstance(child.tag, str))
    rules = []
    for rule in root.xpath("./constraints/rule"):
        expression = next((child for child in rule if isinstance(child.tag, str)), None)
        if expression is not None:
            rules.append(f"<li>{html.escape(_constraint(expression))}</li>")
    corpus = item["corpus"].capitalize()
    return f'''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>{html.escape(display_code)}</title>
<style>body{{max-width:1100px;margin:2rem auto;padding:0 1rem;font:16px/1.5 system-ui,sans-serif;color:#17202a}}h1,h2{{color:#17365d}}ul{{margin:.15rem 0 .5rem 1.5rem;padding-left:1rem}}li{{margin:.35rem 0}}small{{color:#526779}}.description{{margin:.1rem 0 .25rem 1rem;color:#354657}}@media print{{body{{max-width:none}}li{{break-inside:avoid}}}}</style>
</head><body><h1>{html.escape(display_code)} — {html.escape(corpus)}</h1>
<p>Standardized feature-tree view. Group labels: AND = compatible children; OR = at least one; ALT = exactly one. “Mandatory” and “optional” qualify a child under its parent. Citation markers and generation metadata are omitted here; citation evidence is rated separately.</p>
<h2>Feature hierarchy</h2><ul>{_render_node(feature_root)}</ul>
<h2>Cross-tree constraints</h2>{'<ol>' + ''.join(rules) + '</ol>' if rules else '<p>None recorded.</p>'}
</body></html>\n'''


def _read_citation_items(path: Path, items: list[dict], selection_dir: Path) -> list[dict]:
    selected = {item["run_id"]: item for item in items}
    all_docs = _doc_index()
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    if len(rows) != 15:
        raise ValueError(f"expected exactly 15 citation items, got {len(rows)}")
    with (selection_dir / "citation_items_to_complete.csv").open(newline="", encoding="utf-8") as handle:
        frozen_rows = list(csv.DictReader(handle))
    identity = lambda row: (row["run_id"], row["feature"], row["doc_id"])
    if sorted(identity(row) for row in rows) != sorted(identity(row) for row in frozen_rows):
        raise ValueError("citation identities differ from the frozen balanced sample")
    seen = set()
    for row in rows:
        run_id = row.get("run_id", "").strip()
        if run_id not in selected:
            raise ValueError(f"citation run_id is not in selected FMs: {run_id}")
        feature, doc_id = row.get("feature", "").strip(), row.get("doc_id", "").strip()
        if (feature, doc_id) not in _feature_citations(Path(selected[run_id]["xml_path"])):
            raise ValueError(f"feature/document pair was not cited in {run_id}: {feature}, {doc_id}")
        if doc_id not in all_docs or not doc_id.startswith("rep_" if selected[run_id]["corpus"] == "repair" else "fed_"):
            raise ValueError(f"invalid cited document for {run_id}: {doc_id}")
        for field in ("excerpt", "excerpt_origin", "excerpt_locator"):
            if not row.get(field, "").strip():
                raise ValueError(f"citation item missing {field}: {run_id} {feature} {doc_id}")
        if row["excerpt_origin"].strip().lower() not in ("retrieved chunk", "document excerpt"):
            raise ValueError(f"excerpt_origin must be 'retrieved chunk' or 'document excerpt': {run_id}")
        if (run_id, feature, doc_id) in seen:
            raise ValueError(f"duplicate citation item: {run_id} {feature} {doc_id}")
        seen.add((run_id, feature, doc_id))
        row["doc_title"] = all_docs[doc_id].get("title", "")
    return rows


def _rater_order(items: list[dict], rng: random.Random, duplicates: dict[str, tuple[str, dict]]) -> list[tuple[str, dict]]:
    triplets = defaultdict(list)
    for item in items:
        triplets[(item["corpus"], item["seed"])].append(item)
    keys = list(triplets)
    for key in keys:
        rng.shuffle(triplets[key])
        chosen = duplicates.get(key[0])
        if chosen and (chosen[1]["corpus"], chosen[1]["seed"]) == key:
            index = next(i for i, item in enumerate(triplets[key]) if item["run_id"] == chosen[1]["run_id"])
            triplets[key][0], triplets[key][index] = triplets[key][index], triplets[key][0]
    ordered = []
    for round_index in range(3):
        rng.shuffle(keys)
        ordered.extend((triplets[key][round_index]["code"], triplets[key][round_index]) for key in keys)
    for corpus in CORPORA:
        duplicate_code, original = duplicates[corpus]
        if original["run_id"] not in {item["run_id"] for _, item in ordered[:8]}:
            raise ValueError(f"duplicate original is not in first third: {corpus}")
        ordered.append((duplicate_code, original))
    # Both repeats are in the final third and far from their originals.
    ordered[-2:] = [ordered[-1], ordered[-2]] if rng.randrange(2) else ordered[-2:]
    return ordered


def pack(args: argparse.Namespace) -> None:
    selection_dir = args.selection_dir.resolve()
    selection = _read_json(selection_dir / "selection.json")
    items = selection["items"]
    if len(items) != 24 or len({item["run_id"] for item in items}) != 24:
        raise ValueError("selection must contain exactly 24 unique completed runs")
    for item in items:
        if _sha256(Path(item["xml_path"])) != item["xml_sha256"]:
            raise ValueError(f"selected XML changed since selection: {item['xml_path']}")
    citation_rows = _read_citation_items(args.citation_items, items, selection_dir)
    output = args.output.resolve()
    if output.exists():
        raise FileExistsError(f"refusing to overwrite packets: {output}")
    rng = random.Random(selection["blind_seed"] + 301)
    used = {item["code"] for item in items}
    duplicates = {}
    for corpus in CORPORA:
        target = rng.choice([item for item in items if item["corpus"] == corpus])
        duplicates[corpus] = (_new_code(rng, used, "FM-"), target)
    lookup = {item["run_id"]: item for item in items}
    source_rows = [
        {"doc_id": doc_id, "title": row.get("title", ""), "doi_url": row.get("doi_url", "")}
        for doc_id, row in sorted(_doc_index().items())
    ]
    output.mkdir(parents=True)
    author = output / "AUTHOR_ONLY"
    author.mkdir()
    key_rows = [{"code": item["code"], "run_id": item["run_id"], "model_key": item["model_key"], "corpus": item["corpus"], "seed": item["seed"], "duplicate_of": ""} for item in items]
    packet_hashes = []
    for n in range(1, 4):
        rater_code = f"R{n:02d}"
        rater_dir = output / rater_code
        model_dir = rater_dir / "models"
        model_dir.mkdir(parents=True)
        per_rater_rng = random.Random(selection["blind_seed"] + n * 101)
        order = _rater_order(items, per_rater_rng, duplicates)
        wb = load_workbook(args.form)
        ws = wb["1. Rate the models"]
        for row_index, (display_code, item) in enumerate(order, 2):
            ws.cell(row_index, 2, display_code)
            ws.cell(row_index, 3, item["corpus"].capitalize())
            (model_dir / f"{display_code}.html").write_text(_render_html(item, display_code), encoding="utf-8")
        triplets = defaultdict(list)
        for item in items:
            triplets[(item["corpus"], item["seed"])].append(item)
        triplet_keys = list(triplets)
        per_rater_rng.shuffle(triplet_keys)
        ws = wb["2. Compare triplets"]
        for row_index, key in enumerate(triplet_keys, 2):
            options = triplets[key][:]
            per_rater_rng.shuffle(options)
            ws.cell(row_index, 2, key[0].capitalize())
            ws.cell(row_index, 3, options[0]["set_code"])
            for col, item in enumerate(options, 4):
                ws.cell(row_index, col, item["code"])
        ws = wb["3. Check citations"]
        shuffled_citations = citation_rows[:]
        per_rater_rng.shuffle(shuffled_citations)
        for row_index, row in enumerate(shuffled_citations, 2):
            item = lookup[row["run_id"]]
            for col, value in enumerate((item["code"], item["corpus"].capitalize(), row["feature"], f"{row['doc_id']} — {row['doc_title']}", row["excerpt"], f"{row['excerpt_origin']}: {row['excerpt_locator']}"), 2):
                ws.cell(row_index, col, value)
        wb["Rater details"]["B3"] = rater_code
        workbook_path = rater_dir / f"Expert-rating-{rater_code}.xlsx"
        wb.save(workbook_path)
        _write_csv(rater_dir / "source-index.csv", ["doc_id", "title", "doi_url"], source_rows)
        (rater_dir / "README.txt").write_text(
            "Read the workbook 'Read me first' tab. Open models/<FM-code>.html for each row. "
            "The study author supplies authorized source-document access separately. "
            "Do not look for model-provider identities or automated scores.\n",
            encoding="utf-8",
        )
        zip_path = output / f"{rater_code}.zip"
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(rater_dir.rglob("*")):
                if path.is_file():
                    archive.write(path, path.relative_to(rater_dir))
        packet_hashes.append({"rater_code": rater_code, "zip_path": str(zip_path), "sha256": _sha256(zip_path)})
    for corpus, (duplicate_code, target) in duplicates.items():
        key_rows.append({"code": duplicate_code, "run_id": target["run_id"], "model_key": target["model_key"], "corpus": corpus, "seed": target["seed"], "duplicate_of": target["code"]})
    _write_csv(author / "code_key.csv", ["code", "run_id", "model_key", "corpus", "seed", "duplicate_of"], key_rows)
    _write_csv(author / "packet_hashes.csv", ["rater_code", "zip_path", "sha256"], packet_hashes)
    (author / "citation_items.json").write_text(json.dumps(citation_rows, indent=2) + "\n", encoding="utf-8")
    print(f"wrote three blinded packets to {output}; AUTHOR_ONLY must never be sent")


def legacy_main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p_select = sub.add_parser("select", help="freeze 24 matched FMs and emit citation candidates")
    p_select.add_argument("--open-inventory", type=Path, required=True, help="runs.json from the frozen open-weight campaign inventory")
    p_select.add_argument("--astra-results", type=Path, required=True, help="root of separate Astra N=1 RAG extension campaign")
    p_select.add_argument("--output", type=Path, required=True, help="new AUTHOR-ONLY selection directory")
    p_select.add_argument("--blind-seed", type=int, default=20260929)
    p_select.set_defaults(func=select)
    p_pack = sub.add_parser("pack", help="create three rater packets from a frozen selection")
    p_pack.add_argument("--selection-dir", type=Path, required=True)
    p_pack.add_argument("--citation-items", type=Path, required=True, help="15 audited feature-document rows completed from citation_candidates.csv")
    p_pack.add_argument("--form", type=Path, default=REPO / "data/Expert-evaluation-form-v2.xlsx")
    p_pack.add_argument("--output", type=Path, required=True, help="new packet directory")
    p_pack.set_defaults(func=pack)
    args = parser.parse_args()
    args.func(args)
    return 0


if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(REPO))
    from scripts.expert_artifacts_v3 import main
    raise SystemExit(main())
