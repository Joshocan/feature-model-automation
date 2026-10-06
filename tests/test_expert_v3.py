import argparse
from collections import Counter
import csv
import json
from pathlib import Path
import random
import zipfile

from openpyxl import load_workbook
import pytest
from scripts import expert_artifacts_v3 as v


def pool():
    return [dict(corpus=c,stratum=s,run_id=f'{c}-{model}',edge=n)
            for c in v.common.CORPORA for s in ['agree','disagree'] for model in range(5) for n in range(4)]


def test_sample_is_reproducible_stratified_and_capped():
    rows=v.sample_edges(pool())
    assert rows==v.sample_edges(pool())
    assert Counter((r['corpus'],r['stratum']) for r in rows)=={(c,s):3 for c in v.common.CORPORA for s in ['agree','disagree']}
    assert max(Counter(r['run_id'] for r in rows).values())<=2
    assert len({r['item_id'] for r in rows})==12


def test_missing_stratum_fails():
    with pytest.raises(ValueError,match='Insufficient'):
        v.sample_edges([r for r in pool() if r['stratum']=='agree'])


def fixtures():
    items=[dict(code=code,corpus=c,run_id=f'r{index}-{c}',model_id='hidden')
           for c,codes in v.CODES.items() for index,code in enumerate(codes)]
    citations=[dict(item_id=f'C{i:02d}',code=items[i%10]['code'],corpus=items[i%10]['corpus'],feature='R / X',doc_id='rep_01',doc_title='Paper',excerpt='Evidence',excerpt_origin='document excerpt',excerpt_locator='p. 1') for i in range(12)]
    return items,citations


def test_v3_workbook_columns_blanks_and_counterbalance():
    items,rows=fixtures();form=v.REPO/'data/Expert-evaluation-form-v3.xlsx'
    for rater,first in [(1,'Repair'),(2,'Federation'),(3,'Repair')]:
        wb=load_workbook(form);order=v.fill_workbook(wb,items,rows,rater)
        assert wb['1. Rate the models']['B2'].value==first
        assert wb['2. Rank within domain']['D1'].value==first
        assert len({wb['1. Rate the models'].cell(r,1).value for r in range(2,12)})==10
        for r in range(2,14):
            assert 'Evidence' in wb['3. Check citations'].cell(r,6).value
            assert wb['3. Check citations'].cell(r,7).value is None
            assert wb['3. Check citations'].cell(r,8).value is None
        assert {wb['2. Rank within domain'].cell(r,3).value for r in list(range(3,8))+list(range(13,18))}=={i['code'] for i in items}


def test_completed_form_rejected():
    items,rows=fixtures();wb=load_workbook(v.REPO/'data/Expert-evaluation-form-v3.xlsx')
    wb['3. Check citations']['G2']='Yes'
    with pytest.raises(ValueError,match='responses'):v.fill_workbook(wb,items,rows,1)


def test_renderer_collapsible_and_trace_removed(tmp_path):
    path=tmp_path/'fm.xml'
    path.write_text('<featureModel><struct><and name="Root"><feature name="Child"><description>Text Trace: [rep_01]</description></feature></and></struct></featureModel>')
    page=v.common._render_html(dict(xml_path=str(path),corpus='repair'),'Rep-96')
    assert '<details open>' in page and 'optional' in page
    assert 'rep_01' not in page and 'Trace:' not in page


def test_pack_roundtrip_and_identity_guard(tmp_path):
    items,citations=fixtures()
    for i in items:
        xml=tmp_path/(i['code']+'.xml')
        xml.write_text('<featureModel><struct><and name="Root"><feature name="Child"/></and></struct></featureModel>')
        i.update(xml_path=str(xml),xml_sha256=v.common._sha256(xml),N=1,seed=0)
    selection=tmp_path/'selection';selection.mkdir()
    (selection/'selection.json').write_text(json.dumps(dict(version=3,items=items)))
    v.freeze_files(selection,['selection.json'])
    citation_dir=tmp_path/'citations';citation_dir.mkdir()
    blank=[]
    for r in citations:
        blank.append(dict(r,**{k:'' for k in v.EDIT_FIELDS}))
        r['source_checked_by']='reviewer-B'
    v.write_csv(citation_dir/'citation_items_to_complete.csv',blank)
    v.write_csv(citation_dir/'citation_key_AUTHOR_ONLY.csv',[dict(item_id=r['item_id'],stratum='agree') for r in citations])
    (citation_dir/'summary.json').write_text(json.dumps(dict(selection_sha256=v.common._sha256(selection/'selection.json'))))
    v.freeze_files(citation_dir,['citation_items_to_complete.csv','citation_key_AUTHOR_ONLY.csv','summary.json'])
    final=tmp_path/'final.csv';v.write_csv(final,citations)
    args=argparse.Namespace(selection_dir=selection,citation_dir=citation_dir,citation_items=final,
        form=v.REPO/'data/Expert-evaluation-form-v3.xlsx',output=tmp_path/'packets',raters=2)
    v.pack(args)
    from io import BytesIO
    for rater in ['R01','R02']:
        with zipfile.ZipFile(args.output/(rater+'.zip')) as z:
            assert len([n for n in z.namelist() if n.startswith('models/')])==10
            assert not any('AUTHOR_ONLY' in n or 'key' in n for n in z.namelist())
            wb=load_workbook(BytesIO(z.read(f'Expert-rating-{rater}.xlsx')))
            assert wb['3. Check citations'].max_row==13
            assert wb['1. Rate the models'].max_row==11
            assert all(c.comment is None for s in wb for row in s for c in row)
            assert all(wb['3. Check citations'].cell(r,7).value is None for r in range(2,14))
    citations[0]['doc_id']='rep_99';v.write_csv(final,citations)
    with pytest.raises(ValueError,match='identity changed'):v.validated_citations(args,dict(items=items))
